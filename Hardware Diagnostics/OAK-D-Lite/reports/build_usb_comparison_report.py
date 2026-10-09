"""Build a traceable USB 2/USB 3 evidence report for the OAK-D Lite.

The module deliberately separates evidence extraction from document layout.  It
never modifies benchmark records and rejects a comparison unless the two runs
represent the verified device, complete profile set, matching provenance, and
the intended HIGH-to-SUPER transition.
"""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

EXPECTED_MXID = "REPLACE_WITH_CAMERA_MXID"
EXPECTED_PRODUCT = "OAK-D-LITE"
BASELINE_RUN_ID = "20260814T024121316782Z_run_b381fce8"
USB3_INVENTORY_ID = "20260814T034056092598Z_inventory_0c04d31e"
CANDIDATE_RUN_ID = "20260814T034134853796Z_run_ff85bad7"
COMPARISON_ID = "20260814T034425097010Z_comparison_55e31788"

EXPECTED_TEST_IDS = {
    "rgb_preview_640x360_30",
    "rgb_video_1280x720_30",
    "rgb_video_1920x1080_30",
    "mono_pair_640x480_30",
    "stereo_depth_640x480_30",
    "h264_rgb_1920x1080_30",
    "combined_h264_depth_system_30",
    "combined_h264_depth_system_30_stability",
}

PROFILE_STREAMS = [
    ("rgb_preview_640x360_30", "rgb", "RGB", "RGB preview 640 x 360 @ 30"),
    ("rgb_video_1280x720_30", "rgb", "RGB", "Raw RGB 1280 x 720 @ 30"),
    ("rgb_video_1920x1080_30", "rgb", "RGB", "Raw RGB 1920 x 1080 @ 30"),
    ("mono_pair_640x480_30", "left", "Left", "Mono pair 640 x 480 @ 30"),
    ("mono_pair_640x480_30", "right", "Right", "Mono pair 640 x 480 @ 30"),
    ("stereo_depth_640x480_30", "depth", "Depth", "Stereo depth 640 x 480 @ 30"),
    ("h264_rgb_1920x1080_30", "h264", "H.264", "H.264 RGB 1920 x 1080 @ 30"),
    (
        "combined_h264_depth_system_30",
        "depth",
        "Depth",
        "Combined H.264 + depth @ 30 (10 s)",
    ),
    (
        "combined_h264_depth_system_30",
        "h264",
        "H.264",
        "Combined H.264 + depth @ 30 (10 s)",
    ),
    (
        "combined_h264_depth_system_30_stability",
        "depth",
        "Depth",
        "Combined stability (60 s)",
    ),
    (
        "combined_h264_depth_system_30_stability",
        "h264",
        "H.264",
        "Combined stability (60 s)",
    ),
]

TOKENS = {
    "page": {
        "width_in": 11.0,
        "height_in": 8.5,
        "margin_top_in": 0.65,
        "margin_right_in": 0.65,
        "margin_bottom_in": 0.65,
        "margin_left_in": 0.65,
        "header_in": 0.32,
        "footer_in": 0.32,
        "content_width_dxa": 13968,
    },
    "body": {"font": "Calibri", "size_pt": 11, "after_pt": 6, "line_spacing": 1.25},
    "title": {"size_pt": 24, "color": "0B2545", "before_pt": 0, "after_pt": 5},
    "h1": {"size_pt": 16, "color": "2E74B5", "before_pt": 18, "after_pt": 10},
    "h2": {"size_pt": 13, "color": "2E74B5", "before_pt": 14, "after_pt": 7},
    "h3": {"size_pt": 12, "color": "1F4D78", "before_pt": 10, "after_pt": 5},
    "table": {
        "font_pt": 8.7,
        "header_font_pt": 8.5,
        "indent_dxa": 120,
        "cell_margins_dxa": {"top": 80, "bottom": 80, "start": 120, "end": 120},
    },
    "colors": {
        "navy": "0B2545",
        "blue": "2E74B5",
        "teal": "167D8D",
        "header_fill": "E8EEF5",
        "alternate_fill": "F4F7FA",
        "border": "AAB7C4",
        "positive_fill": "E8F3F1",
        "pending_fill": "FFF6DD",
        "ambiguous_fill": "FDECEC",
        "muted": "5B6573",
        "white": "FFFFFF",
    },
}

TABLE_WIDTHS = {
    "masthead": [3492, 3492, 3492, 3492],
    "context": [2400, 4584, 2400, 4584],
    "specifications": [3000, 3600, 4560, 2808],
    "comparison": [3000, 936, 936, 720, 1080, 936, 936, 720, 1080, 936, 2688],
    "stability": [2600, 1100, 1500, 1300, 1700, 1800, 3968],
    "constraints": [2300, 3600, 3400, 1500, 3168],
    "evidence": [3300, 10668],
}


def _reject_non_finite(value: str) -> None:
    raise ValueError(f"non-standard JSON constant: {value}")


def load_strict_json(path: Path) -> dict[str, Any]:
    """Load one evidence document and reject NaN/Infinity tokens."""

    with Path(path).open(encoding="utf-8") as handle:
        document = json.load(handle, parse_constant=_reject_non_finite)

    if not isinstance(document, dict):
        raise ValueError(f"evidence root must be an object: {path}")

    return document


def metric_value(record: dict[str, Any], key: str) -> Any:
    """Return the value field for a named metric record."""

    metric = record[key]

    if not isinstance(metric, dict) or "value" not in metric:
        raise ValueError(f"metric {key!r} is not a value record")

    return metric["value"]


def _metric(record: dict[str, Any]) -> Any:
    if not isinstance(record, dict) or "value" not in record:
        raise ValueError("expected metric record with a value field")

    return record["value"]


def _stream_values(
    run: dict[str, Any], profile_id: str, stream_id: str
) -> dict[str, Any]:
    # Extract the six frame-delivery metrics shown for one profile stream.
    stream = run["tests"][profile_id]["streams"][stream_id]
    metrics = stream["metrics"]
    return {
        "fps": metric_value(metrics, "received_fps"),
        "frames": metric_value(metrics, "frames_received"),
        "gaps": metric_value(metrics, "sequence_gaps"),
        "payload_mbps": metric_value(metrics, "payload_bitrate_mbps"),
        "p95_ms": metric_value(metrics, "latency_p95_ms"),
        "jitter_ms": metric_value(metrics, "interarrival_jitter_ms"),
    }


def _percent_change(baseline: float, candidate: float) -> float | None:
    # Return percentage change relative to the baseline, or None for a zero baseline.
    if baseline == 0:
        return None

    return (candidate - baseline) / baseline * 100.0


def _observation(baseline: dict[str, Any], candidate: dict[str, Any]) -> str:
    # Describe the FPS, sequence-gap and p95-latency changes for a stream pair.
    fps_change = _percent_change(baseline["fps"], candidate["fps"])
    p95_change = _percent_change(baseline["p95_ms"], candidate["p95_ms"])
    parts = []

    if fps_change is not None:
        parts.append(f"FPS {fps_change:+.1f}%")

    if baseline["gaps"] > 0 and candidate["gaps"] == 0:
        parts.append(f"gaps {baseline['gaps']} to 0")

    elif baseline["gaps"] != candidate["gaps"]:
        parts.append(f"gaps {baseline['gaps']} to {candidate['gaps']}")

    else:
        parts.append(f"gaps unchanged at {candidate['gaps']}")

    if p95_change is not None:
        parts.append(f"p95 {p95_change:+.1f}%")

    return "; ".join(parts)


def _specification_rows(candidate: dict[str, Any]) -> list[dict[str, str]]:
    # Combine candidate inventory, calibration and labelled claims into specification
    # rows.
    claims = candidate["manufacturer_claims"]
    calibration = candidate["calibration"]
    eeprom = _metric(calibration["eeprom_json"])
    cameras = {item["socket"]: item for item in candidate["device"]["camera_features"]}
    fovs = calibration["cameras"]
    return [
        {
            "parameter": "Product",
            "value": _metric(candidate["device"]["product_name"]),
            "evidence": "DepthAI device inventory",
            "status": "Observed",
        },
        {
            "parameter": "Board",
            "value": f"{eeprom['boardName']}, revision {eeprom['boardRev']}",
            "evidence": "Calibration EEPROM",
            "status": "Calibrated",
        },
        {
            "parameter": "Dimensions",
            "value": "91 x 28 x 17.5 mm",
            "evidence": claims["dimensions"]["method"],
            "status": "Claimed",
        },
        {
            "parameter": "RGB camera",
            "value": (
                f"CAM_A {cameras['CAM_A']['sensor_name']}; "
                f"{cameras['CAM_A']['width']} x {cameras['CAM_A']['height']}; "
                "colour; autofocus; claimed maximum 60 FPS"
            ),
            "evidence": "DepthAI inventory plus supplied datasheet",
            "status": "Observed + claimed",
        },
        {
            "parameter": "Stereo mono cameras",
            "value": (
                f"CAM_B/C {cameras['CAM_B']['sensor_name']}; "
                f"{cameras['CAM_B']['width']} x {cameras['CAM_B']['height']}; "
                "mono global shutter; fixed focus; claimed maximum 120 FPS"
            ),
            "evidence": "DepthAI inventory plus supplied datasheet",
            "status": "Observed + claimed",
        },
        {
            "parameter": "Stereo baseline",
            "value": f"{_metric(claims['stereo_baseline'])} mm",
            "evidence": claims["stereo_baseline"]["method"],
            "status": "Claimed",
        },
        {
            "parameter": "Horizontal field of view",
            "value": (
                f"calibrated CAM_A {metric_value(fovs['CAM_A'], 'horizontal_fov_calculated'):.3f} deg; "
                f"CAM_B {metric_value(fovs['CAM_B'], 'horizontal_fov_calculated'):.3f} deg; "
                f"CAM_C {metric_value(fovs['CAM_C'], 'horizontal_fov_calculated'):.3f} deg"
            ),
            "evidence": "DepthAI calibration handler",
            "status": "Calibrated",
        },
        {
            "parameter": "USB interface wording",
            "value": "USB 3.1 Gen1 and USB 3.1 Gen2 both appear in the supplied datasheet",
            "evidence": claims["usb_interface"]["method"],
            "status": "Claimed ‡",
        },
        {
            "parameter": "USB power wording",
            "value": _metric(claims["usb2_max_wording"]),
            "evidence": claims["usb2_max_wording"]["method"],
            "status": "Claimed ‡",
        },
        {
            "parameter": "Input and power",
            "value": "recommended 5-5.25 V; maximum 1.5 A; claimed 4-6 W operating / 2.5 W idle",
            "evidence": "Supplied manufacturer datasheet",
            "status": "Claimed",
        },
        {
            "parameter": "Ambient limit",
            "value": f"{_metric(claims['ambient_max'])} deg C ambient",
            "evidence": claims["ambient_max"]["method"],
            "status": "Claimed; not chip temperature",
        },
        {
            "parameter": "Connected IMU",
            "value": candidate["device"]["connected_imu"],
            "evidence": "DepthAI device inventory",
            "status": "Observed",
        },
        {
            "parameter": "Runtime USB link",
            "value": "USB 2 baseline HIGH; USB 3 candidate SUPER",
            "evidence": "DepthAI Device.getUsbSpeed after device boot",
            "status": "Observed *",
        },
        {
            "parameter": "DepthAI",
            "value": candidate["software"]["depthai_version"],
            "evidence": "Python runtime package version",
            "status": "Observed",
        },
    ]


def _stability_rows(
    baseline: dict[str, Any], candidate: dict[str, Any]
) -> list[dict[str, Any]]:
    # Compare delivery, temperature and CPU readings from the two 60-second runs.
    profile = "combined_h264_depth_system_30_stability"
    baseline_system = baseline["tests"][profile]["system"]["metrics"]
    candidate_system = candidate["tests"][profile]["system"]["metrics"]
    rows: list[dict[str, Any]] = []

    for stream_id, label in (("depth", "Depth stream"), ("h264", "H.264 stream")):
        usb2 = _stream_values(baseline, profile, stream_id)
        usb3 = _stream_values(candidate, profile, stream_id)
        rows.extend(
            [
                {
                    "measurement": f"{label} received FPS",
                    "usb2": usb2["fps"],
                    "usb3": usb3["fps"],
                    "unit": "frames/s",
                    "status": "Measured *",
                },
                {
                    "measurement": f"{label} sequence gaps",
                    "usb2": usb2["gaps"],
                    "usb3": usb3["gaps"],
                    "unit": "frames",
                    "status": "Measured *",
                },
                {
                    "measurement": f"{label} p95 latency",
                    "usb2": usb2["p95_ms"],
                    "usb3": usb3["p95_ms"],
                    "unit": "ms",
                    "status": "Measured *",
                },
            ]
        )

    rows.extend(
        [
            {
                "measurement": "Peak internal chip temperature",
                "usb2": max(
                    metric_value(component, "max_c")
                    for component in baseline_system["temperature_c"].values()
                ),
                "usb3": max(
                    metric_value(component, "max_c")
                    for component in candidate_system["temperature_c"].values()
                ),
                "unit": "deg C internal",
                "status": "Measured *",
            },
            {
                "measurement": "LEON CSS peak utilisation",
                "usb2": metric_value(
                    baseline_system["leon_cpu_fraction"]["css"],
                    "peak",
                )
                * 100.0,
                "usb3": metric_value(
                    candidate_system["leon_cpu_fraction"]["css"],
                    "peak",
                )
                * 100.0,
                "unit": "%",
                "status": "Measured *",
            },
            {
                "measurement": "LEON MSS peak utilisation",
                "usb2": metric_value(
                    baseline_system["leon_cpu_fraction"]["mss"],
                    "peak",
                )
                * 100.0,
                "usb3": metric_value(
                    candidate_system["leon_cpu_fraction"]["mss"],
                    "peak",
                )
                * 100.0,
                "unit": "%",
                "status": "Measured *",
            },
        ]
    )
    return rows


def _constraint_rows() -> list[dict[str, str]]:
    # Return the fixed constraints and follow-up wording for these archived benchmark
    # runs.
    return [
        {
            "id": "CON-USB-01",
            "constraint": "Runtime link negotiation",
            "evidence": "The same OAK changed from HIGH to SUPER when the cable/link changed.",
            "consequence": "Record getUsbSpeed() after boot; cable labels and port capability alone are insufficient evidence.",
            "status": "Measured *",
            "follow_up": "Confirm SUPER on the Raspberry Pi before deployment tests.",
        },
        {
            "id": "CON-RAW-02",
            "constraint": "Raw RGB link demand",
            "evidence": "USB 2 raw 720p had 21 gaps; raw 1080p delivered 12.960 FPS with 156 gaps. USB 3 delivered both near 30 FPS with zero gaps.",
            "consequence": "Raw high-resolution transfer needs USB 3 headroom in the tested architecture.",
            "status": "Measured *",
            "follow_up": "Repeat raw profiles on the Raspberry Pi using the validated cable.",
        },
        {
            "id": "CON-ENC-03",
            "constraint": "Encoded-video bandwidth",
            "evidence": "OAK-encoded 1080p H.264 remained near 30 FPS at about 8.5 Mbit/s on both links with zero gaps.",
            "consequence": "On-device encoding substantially reduces the host-link payload compared with raw 1080p.",
            "status": "Measured *",
            "follow_up": "Measure decode, overlay and web delivery on the deployed system.",
        },
        {
            "id": "CON-DEP-04",
            "constraint": "Depth output bandwidth",
            "evidence": "640 x 480 depth used about 148 Mbit/s and remained near 30 FPS on both tested links.",
            "consequence": "Depth is a material concurrent stream even when raw mono inputs remain on the OAK.",
            "status": "Measured *",
            "follow_up": "Include depth in Pi/network resource and data-budget tests.",
        },
        {
            "id": "CON-STB-05",
            "constraint": "Concurrent short-term stability",
            "evidence": "H.264, depth and telemetry completed for 60 seconds with both streams near 30 FPS and zero gaps; peak internal temperature stayed below 45 deg C.",
            "consequence": "The combined OAK workload was stable for the tested duration, not for the full acceptance duration.",
            "status": "Measured *",
            "follow_up": "Repeat on Pi and extend to the required logged operating duration.",
        },
        {
            "id": "CON-PI-06",
            "constraint": "Deployment-host dependence",
            "evidence": "Both controlled runs used a MacBook Pro M4 Pro through the same USB3.1 hub.",
            "consequence": "These results do not establish Raspberry Pi USB, CPU, power or scheduling performance.",
            "status": "Pending †",
            "follow_up": "Run inventory and the full profile set on the Raspberry Pi 5.",
        },
        {
            "id": "CON-MET-07",
            "constraint": "Payload measurement boundary",
            "evidence": "Reported bitrate counts received application payload only; USB framing, XLink overhead and unobserved packets are excluded.",
            "consequence": "Do not use the payload bitrate as total USB-bus utilisation.",
            "status": "Method limit *",
            "follow_up": "Use a bus analyser only if total-bus utilisation becomes necessary.",
        },
        {
            "id": "CON-NET-08",
            "constraint": "Network and browser delivery",
            "evidence": "The benchmark ends at host receipt and does not transmit to the QUT network or browser.",
            "consequence": "REQ-F-04, REQ-F-07 and the four-second REQ-M-19 path remain unverified.",
            "status": "Pending †",
            "follow_up": "Measure capture-to-browser latency, loss and reconnect behaviour.",
        },
        {
            "id": "CON-OVR-09",
            "constraint": "Overlay processing",
            "evidence": "No overlay drawing, compositing or encoded-overlay path was exercised.",
            "consequence": "Overlay clarity, alignment, update rate and compute cost are unknown.",
            "status": "Pending †",
            "follow_up": "Define overlay content and measure its frame-rate and latency impact.",
        },
        {
            "id": "CON-TGT-10",
            "constraint": "Target-analysis performance",
            "evidence": "No 200 mm ArUco, pressure-gauge or valve scene was evaluated.",
            "consequence": "Detection range, pose accuracy, reading error and classification performance cannot be inferred.",
            "status": "Pending †",
            "follow_up": "Create scene-specific acceptance criteria and datasets.",
        },
        {
            "id": "CON-OPS-11",
            "constraint": "Operating-scene dependence",
            "evidence": "The static bench test did not reproduce 1-3 m altitude, up to 2 m/s motion, changing light or target orientations.",
            "consequence": "Bench delivery results do not establish performance in the stated operating context.",
            "status": "Pending †",
            "follow_up": "Test representative distance, motion, lighting and target orientation.",
        },
        {
            "id": "CON-ONB-12",
            "constraint": "Onboard-processing scope",
            "evidence": "StereoDepth, H.264 and SystemLogger ran on the OAK; target-analysis placement was not tested.",
            "consequence": "The results partially support onboard processing but do not allocate all computer-vision work between OAK and Pi.",
            "status": "Partial †",
            "follow_up": "Profile the chosen detector/ArUco pipeline on its intended processor.",
        },
        {
            "id": "CON-DUR-13",
            "constraint": "Acceptance-duration evidence",
            "evidence": "The longest combined benchmark execution was 60 seconds; REQ-M-15 calls for at least 10 minutes of logged operation.",
            "consequence": "The present stability result is engineering evidence, not duration acceptance.",
            "status": "Pending †",
            "follow_up": "Run a logged 10-minute or longer deployed-system test.",
        },
    ]


def build_report_model(root: Path) -> dict[str, Any]:
    """Build and validate the presentation model from immutable evidence."""

    root = Path(root)
    results = root / "results"

    # Load the archived USB 2 run, USB 3 inventory/run and their saved comparison.
    paths = {
        "baseline": results / BASELINE_RUN_ID / "raw.json",
        "inventory": results / USB3_INVENTORY_ID / "raw.json",
        "candidate": results / CANDIDATE_RUN_ID / "raw.json",
        "comparison": results / COMPARISON_ID / "comparison.json",
    }
    baseline = load_strict_json(paths["baseline"])
    inventory = load_strict_json(paths["inventory"])
    candidate = load_strict_json(paths["candidate"])
    comparison = load_strict_json(paths["comparison"])
    comparison_meta = comparison["comparison"]["_comparison"]

    comparison_rows = []

    # Pair each expected stream's USB 2/USB 3 metrics with a summary of its changes.
    for profile_id, stream_id, stream_label, display_name in PROFILE_STREAMS:
        usb2 = _stream_values(baseline, profile_id, stream_id)
        usb3 = _stream_values(candidate, profile_id, stream_id)
        comparison_rows.append(
            {
                "profile_id": profile_id,
                "profile": display_name,
                "stream_id": stream_id,
                "stream_label": stream_label,
                "usb2": usb2,
                "usb3": usb3,
                "observation": _observation(usb2, usb3),
                "status": "Measured *",
            }
        )

    # Collect the three recorded MXIDs and mark whether they identify the same unit.
    baseline_mxid = _metric(baseline["device"]["mxid"])
    candidate_mxid = _metric(candidate["device"]["mxid"])
    inventory_mxid = _metric(inventory["device"]["mxid"])
    identity = {
        "mxid": candidate_mxid,
        "product_name": _metric(candidate["device"]["product_name"]),
        "device_name": _metric(candidate["device"]["device_name"]),
        "board_name": _metric(candidate["calibration"]["eeprom_json"])["boardName"],
        "board_revision": _metric(candidate["calibration"]["eeprom_json"])["boardRev"],
        "all_mxids_match": len({baseline_mxid, candidate_mxid, inventory_mxid}) == 1,
    }
    model = {
        "identity": identity,
        "links": {
            "usb2": _metric(baseline["device"]["usb_speed"]),
            "usb3": _metric(candidate["device"]["usb_speed"]),
        },
        "test_context": {
            "host": candidate["operator_conditions"]["host_label"],
            "hub": candidate["operator_conditions"]["hub"],
            "port": candidate["operator_conditions"]["port"],
            "power": candidate["operator_conditions"]["power"],
            "baseline_cable": baseline["operator_conditions"]["cable"],
            "candidate_cable": candidate["operator_conditions"]["cable"],
            "depthai_version": candidate["software"]["depthai_version"],
            "schema_version": candidate["schema_version"],
            "standard_duration_s": 10,
            "stability_duration_s": 60,
            "baseline_timestamp": baseline["timestamps"]["local"],
            "candidate_timestamp": candidate["timestamps"]["local"],
        },
        "specification_rows": _specification_rows(candidate),
        "comparison_rows": comparison_rows,
        "stability_rows": _stability_rows(baseline, candidate),
        "constraint_rows": _constraint_rows(),
        "evidence_references": [
            {
                "label": "USB 2 baseline",
                "id": BASELINE_RUN_ID,
                "path": str(paths["baseline"].relative_to(root)),
            },
            {
                "label": "USB 3 inventory",
                "id": USB3_INVENTORY_ID,
                "path": str(paths["inventory"].relative_to(root)),
            },
            {
                "label": "USB 3 candidate",
                "id": CANDIDATE_RUN_ID,
                "path": str(paths["candidate"].relative_to(root)),
            },
            {
                "label": "Generated comparison",
                "id": COMPARISON_ID,
                "path": str(paths["comparison"].relative_to(root)),
            },
        ],
        "integrity": {
            "baseline_status": baseline["status"],
            "candidate_status": candidate["status"],
            "inventory_status": inventory["status"],
            "baseline_test_ids": sorted(baseline["tests"]),
            "candidate_test_ids": sorted(candidate["tests"]),
            "baseline_complete_execution_count": sum(
                test["status"] == "complete" for test in baseline["tests"].values()
            ),
            "complete_execution_count": sum(
                test["status"] == "complete" for test in candidate["tests"].values()
            ),
            "provenance_match": baseline["benchmark_provenance"]
            == candidate["benchmark_provenance"]
            == inventory["benchmark_provenance"],
            "benchmark_provenance": candidate["benchmark_provenance"],
            "comparison_status": comparison_meta["status"],
            "comparison_warnings": comparison_meta["warnings"],
            "comparison_provenance_match": comparison_meta["provenance_match"],
            "comparison_usb2": comparison_meta["baseline_usb_speed"],
            "comparison_usb3": comparison_meta["candidate_usb_speed"],
        },
        "warnings": sorted(
            set(baseline.get("warnings", []) + candidate.get("warnings", []))
        ),
    }

    # Validate the recorded evidence before any document layout is built.
    validate_report_model(model)
    return model


def validate_report_model(model: dict[str, Any]) -> None:
    """Reject a model that cannot support a controlled comparison report."""

    identity = model["identity"]
    integrity = model["integrity"]

    # Reject a different camera or inconsistent unit IDs before accepting any
    # measurements.
    if (
        identity["mxid"] != EXPECTED_MXID
        or not identity["all_mxids_match"]
    ):
        raise ValueError("evidence MXID does not match the verified OAK-D Lite")

    if identity["product_name"] != EXPECTED_PRODUCT:
        raise ValueError("evidence product is not OAK-D-LITE")

    if model["links"] != {"usb2": "HIGH", "usb3": "SUPER"}:
        raise ValueError("evidence is not the intended HIGH-to-SUPER USB transition")

    if (
        integrity["baseline_status"] != "complete"
        or integrity["candidate_status"] != "complete"
    ):
        raise ValueError("both benchmark runs must be complete")

    if integrity["inventory_status"] != "complete":
        raise ValueError("USB 3 inventory must be complete")

    # Require all seven profiles and the stability repeat to be complete in both runs.
    if (
        integrity["baseline_complete_execution_count"] != 8
        or integrity["complete_execution_count"] != 8
    ):
        raise ValueError("comparison requires eight complete executions in each run")

    if (
        set(integrity["baseline_test_ids"]) != EXPECTED_TEST_IDS
        or set(integrity["candidate_test_ids"]) != EXPECTED_TEST_IDS
    ):
        raise ValueError(
            "benchmark profile set does not match the approved eight executions"
        )

    # Reject differing recorded benchmark hashes, even if the measured values look
    # usable.
    if (
        not integrity["provenance_match"]
        or not integrity["comparison_provenance_match"]
    ):
        raise ValueError("benchmark provenance does not match")

    if integrity["comparison_status"] != "expected_usb2_to_usb3_retest":
        raise ValueError("comparison is not classified as the expected USB retest")

    if integrity["comparison_warnings"]:
        raise ValueError("comparison contains warnings")

    if (
        integrity["comparison_usb2"] != "HIGH"
        or integrity["comparison_usb3"] != "SUPER"
    ):
        raise ValueError("comparison metadata does not record HIGH-to-SUPER")

    if len(model["comparison_rows"]) != len(PROFILE_STREAMS):
        raise ValueError("comparison stream row count is incomplete")

    # Reject nonnumeric timings/rates, missing frames or negative gaps in either run.
    for row in model["comparison_rows"]:
        for link in ("usb2", "usb3"):
            values = row[link]

            for key in ("fps", "payload_mbps", "p95_ms", "jitter_ms"):
                if not isinstance(values[key], (int, float)) or not math.isfinite(
                    values[key]
                ):
                    raise ValueError(f"{row['profile_id']} {link} {key} is not finite")

            if values["frames"] <= 0 or values["gaps"] < 0:
                raise ValueError(
                    f"{row['profile_id']} {link} frame evidence is invalid"
                )


def _ensure_child(parent: Any, tag: str) -> Any:
    # Reuse an existing Word XML property or create it when the property is absent.
    child = parent.find(qn(tag))

    if child is None:
        child = OxmlElement(tag)
        parent.append(child)

    return child


def _set_run_font(
    run: Any,
    *,
    size_pt: float,
    color: str = "0B2545",
    bold: bool = False,
    italic: bool = False,
) -> None:
    # Apply the configured font, point size, colour and emphasis to one text run.
    run.font.name = TOKENS["body"]["font"]
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), TOKENS["body"]["font"])
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), TOKENS["body"]["font"])
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def _set_cell_fill(cell: Any, fill: str) -> None:
    # Set the cell's Word shading property to the requested background colour.
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = _ensure_child(tc_pr, "w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_table_borders(table: Any, color: str) -> None:
    # Apply the same line width and colour to outer borders and internal grid lines.
    borders = _ensure_child(table._tbl.tblPr, "w:tblBorders")

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = _ensure_child(borders, f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _set_repeat_header(row: Any) -> None:
    # Mark this row as a header that Word repeats when the table crosses a page.
    tr_pr = row._tr.get_or_add_trPr()
    marker = _ensure_child(tr_pr, "w:tblHeader")
    marker.set(qn("w:val"), "true")


def _prevent_row_split(row: Any) -> None:
    # Keep the row on one page and allow its height to follow its contents.
    tr_pr = row._tr.get_or_add_trPr()
    _ensure_child(tr_pr, "w:cantSplit")
    row.height = None


def _apply_table_geometry(table: Any, widths: list[int]) -> None:
    # Reject column widths that do not fill the configured report width in twips.
    total = TOKENS["page"]["content_width_dxa"]

    if sum(widths) != total:
        raise ValueError(f"table column widths sum to {sum(widths)}, expected {total}")

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    table_width = _ensure_child(tbl_pr, "w:tblW")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total))
    table_indent = _ensure_child(tbl_pr, "w:tblInd")
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(TOKENS["table"]["indent_dxa"]))
    layout = _ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    # Replace the table grid with the requested column widths used by each cell.
    grid = table._tbl.tblGrid

    for child in list(grid):
        grid.remove(child)

    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    margins = TOKENS["table"]["cell_margins_dxa"]

    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)

    for row in table.rows:
        _prevent_row_split(row)

        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            cell_width = _ensure_child(tc_pr, "w:tcW")
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))
            cell_margins = _ensure_child(tc_pr, "w:tcMar")

            for side in ("top", "bottom", "start", "end"):
                margin = _ensure_child(cell_margins, f"w:{side}")
                margin.set(qn("w:type"), "dxa")
                margin.set(qn("w:w"), str(margins[side]))

    _set_table_borders(table, TOKENS["colors"]["border"])


def _format_cell(
    cell: Any,
    *,
    header: bool = False,
    fill: str | None = None,
    align: Any = WD_ALIGN_PARAGRAPH.LEFT,
    font_size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
) -> None:
    # Apply optional shading and format all text runs using header or body defaults.
    if fill:
        _set_cell_fill(cell, fill)

    resolved_size = font_size or (
        TOKENS["table"]["header_font_pt"] if header else TOKENS["table"]["font_pt"]
    )
    resolved_color = color or (
        TOKENS["colors"]["white"] if header else TOKENS["colors"]["navy"]
    )

    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        for run in paragraph.runs:
            _set_run_font(
                run,
                size_pt=resolved_size,
                color=resolved_color,
                bold=header if bold is None else bold,
            )


def _status_fill(status: str) -> str | None:
    # Choose ambiguous, pending or measured shading from the status text.
    if "‡" in status:
        return TOKENS["colors"]["ambiguous_fill"]

    if "†" in status:
        return TOKENS["colors"]["pending_fill"]

    if "Measured" in status or "Observed *" in status:
        return TOKENS["colors"]["positive_fill"]

    return None


def _style_table(table: Any, *, header_rows: int = 1, zebra: bool = True) -> None:
    # Mark header rows to repeat on each page and shade alternating data rows.
    for row_index, row in enumerate(table.rows):
        if row_index < header_rows:
            _set_repeat_header(row)

            for cell in row.cells:
                _format_cell(
                    cell,
                    header=True,
                    fill=TOKENS["colors"]["navy"],
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                )

        else:
            fill = (
                TOKENS["colors"]["alternate_fill"]
                if zebra and row_index % 2 == 0
                else None
            )

            for cell in row.cells:
                _format_cell(cell, fill=fill)


def _set_table_density(
    table: Any, *, font_size: float, vertical_margin_dxa: int
) -> None:
    """Tighten a dense register without changing its semantic styling."""

    # Reduce top/bottom cell margins and set the font size across the table.
    for row in table.rows:
        for cell in row.cells:
            cell_margins = _ensure_child(cell._tc.get_or_add_tcPr(), "w:tcMar")

            for side in ("top", "bottom"):
                margin = _ensure_child(cell_margins, f"w:{side}")
                margin.set(qn("w:type"), "dxa")
                margin.set(qn("w:w"), str(vertical_margin_dxa))

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

                for run in paragraph.runs:
                    run.font.size = Pt(font_size)


def _write_cell(cell: Any, value: Any) -> None:
    cell.text = "" if value is None else str(value)


def _add_standard_table(
    document: DocumentType,
    headers: list[str],
    rows: list[list[Any]],
    widths_key: str,
    *,
    status_column: int | None = None,
    numeric_columns: tuple[int, ...] = (),
) -> Any:
    # Fill a header and data rows, then apply widths, alignment and status shading.
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        _write_cell(table.rows[0].cells[index], header)

    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            _write_cell(table.rows[row_index].cells[column_index], value)

    # Apply fixed widths after filling the cells, then align numeric and status columns.
    _apply_table_geometry(table, TABLE_WIDTHS[widths_key])
    _style_table(table)

    for row_index in range(1, len(table.rows)):
        for column_index in numeric_columns:
            _format_cell(
                table.rows[row_index].cells[column_index],
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )

        if status_column is not None:
            cell = table.rows[row_index].cells[status_column]
            _format_cell(
                cell,
                fill=_status_fill(cell.text),
                align=WD_ALIGN_PARAGRAPH.CENTER,
                bold=True,
            )

    return table


def _set_style(
    style: Any,
    *,
    size_pt: float,
    color: str,
    before_pt: float,
    after_pt: float,
    bold: bool,
    line_spacing: float,
) -> None:
    # Set the named document style's font, colour and paragraph spacing.
    style.font.name = TOKENS["body"]["font"]
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), TOKENS["body"]["font"])
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), TOKENS["body"]["font"])
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(before_pt)
    paragraph.space_after = Pt(after_pt)
    paragraph.line_spacing = line_spacing


def _configure_document(document: DocumentType) -> None:
    # Configure landscape pages, report styles and document metadata from the constants.
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(TOKENS["page"]["width_in"])
    section.page_height = Inches(TOKENS["page"]["height_in"])
    section.top_margin = Inches(TOKENS["page"]["margin_top_in"])
    section.right_margin = Inches(TOKENS["page"]["margin_right_in"])
    section.bottom_margin = Inches(TOKENS["page"]["margin_bottom_in"])
    section.left_margin = Inches(TOKENS["page"]["margin_left_in"])
    section.header_distance = Inches(TOKENS["page"]["header_in"])
    section.footer_distance = Inches(TOKENS["page"]["footer_in"])

    _set_style(
        document.styles["Normal"],
        size_pt=TOKENS["body"]["size_pt"],
        color=TOKENS["colors"]["navy"],
        before_pt=0,
        after_pt=TOKENS["body"]["after_pt"],
        bold=False,
        line_spacing=TOKENS["body"]["line_spacing"],
    )

    for style_name, token_name in (
        ("Heading 1", "h1"),
        ("Heading 2", "h2"),
        ("Heading 3", "h3"),
    ):
        token = TOKENS[token_name]
        _set_style(
            document.styles[style_name],
            size_pt=token["size_pt"],
            color=token["color"],
            before_pt=token["before_pt"],
            after_pt=token["after_pt"],
            bold=True,
            line_spacing=1.0,
        )
        document.styles[style_name].paragraph_format.keep_with_next = True
        document.styles[style_name].paragraph_format.keep_together = True

    properties = document.core_properties
    properties.title = "OAK-D Lite USB 2 / USB 3 Benchmark Comparison"
    properties.subject = "EGH455 hardware evidence and constraints"
    properties.author = "EGH455 Group 17"
    properties.keywords = "OAK-D Lite, USB 2, USB 3, benchmark, constraints"


def _add_page_field(paragraph: Any, field_name: str) -> None:
    # Build the Word field elements needed to display a page number or total page count.
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field_name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    # Insert the PAGE or NUMPAGES field into the footer with an initial displayed value.
    run = paragraph.add_run()
    run._r.extend((begin, instruction, separate, result, end))
    _set_run_font(run, size_pt=8.5, color=TOKENS["colors"]["muted"])


def _add_running_furniture(document: DocumentType) -> None:
    # Add the repeated report header and the Page X of Y footer.
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "EGH455 | OAK-D LITE HARDWARE EVIDENCE\tUSB 2 / USB 3"
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(9.7),
        WD_ALIGN_PARAGRAPH.RIGHT,
    )

    for run in paragraph.runs:
        _set_run_font(
            run,
            size_pt=8.5,
            color=TOKENS["colors"]["muted"],
            bold=True,
        )

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    prefix = paragraph.add_run("Page ")
    _set_run_font(prefix, size_pt=8.5, color=TOKENS["colors"]["muted"])
    _add_page_field(paragraph, "PAGE")
    middle = paragraph.add_run(" of ")
    _set_run_font(middle, size_pt=8.5, color=TOKENS["colors"]["muted"])
    _add_page_field(paragraph, "NUMPAGES")


def _add_title_block(document: DocumentType, model: dict[str, Any]) -> None:
    # Add the report title and four summary cells for identity, links and evidence
    # checks.
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(3)
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run("EGH455 | HARDWARE EVIDENCE")
    _set_run_font(
        run,
        size_pt=9.5,
        color=TOKENS["colors"]["teal"],
        bold=True,
    )

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(TOKENS["title"]["before_pt"])
    title.paragraph_format.space_after = Pt(TOKENS["title"]["after_pt"])
    run = title.add_run("OAK-D Lite USB 2 / USB 3 Benchmark Comparison")
    _set_run_font(
        run,
        size_pt=TOKENS["title"]["size_pt"],
        color=TOKENS["title"]["color"],
        bold=True,
    )

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run(
        "Verified hardware specifications, stream measurements and evidence-backed constraints"
    )
    _set_run_font(run, size_pt=11.5, color=TOKENS["colors"]["muted"])

    cells = [
        ("DEVICE", f"OAK-D Lite\n{model['identity']['mxid']}"),
        ("RUNTIME LINK", "HIGH  ->  SUPER"),
        ("EXECUTIONS", "8 / 8 complete\nfor each link"),
        ("PROVENANCE", "Exact match\n0 comparison warnings"),
    ]
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    for index, (label, value) in enumerate(cells):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = paragraph.add_run(label + "\n")
        _set_run_font(
            label_run,
            size_pt=7.8,
            color=TOKENS["colors"]["teal"],
            bold=True,
        )
        value_run = paragraph.add_run(value)
        _set_run_font(
            value_run,
            size_pt=9.5,
            color=TOKENS["colors"]["navy"],
            bold=True,
        )
        _set_cell_fill(cell, TOKENS["colors"]["alternate_fill"])

    _apply_table_geometry(table, TABLE_WIDTHS["masthead"])

    for cell in table.rows[0].cells:
        _format_cell(
            cell,
            fill=TOKENS["colors"]["alternate_fill"],
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    _set_repeat_header(table.rows[0])


def _add_section_heading(document: DocumentType, text: str) -> None:
    document.add_paragraph(text, style="Heading 1")


def _add_note(
    document: DocumentType, label: str, text: str, *, fill: str = "E8EEF5"
) -> None:
    # Add a shaded paragraph with a coloured border, bold label and explanatory text.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(8)
    paragraph.paragraph_format.right_indent = Pt(8)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph_format = paragraph._p.get_or_add_pPr()
    shading = _ensure_child(paragraph_format, "w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    borders = _ensure_child(paragraph_format, "w:pBdr")
    left_border = _ensure_child(borders, "w:left")
    left_border.set(qn("w:val"), "single")
    left_border.set(qn("w:sz"), "18")
    left_border.set(qn("w:space"), "5")
    left_border.set(qn("w:color"), TOKENS["colors"]["teal"])
    label_run = paragraph.add_run(label + "  ")
    _set_run_font(
        label_run,
        size_pt=9.2,
        color=TOKENS["colors"]["teal"],
        bold=True,
    )
    text_run = paragraph.add_run(text)
    _set_run_font(text_run, size_pt=9.2, color=TOKENS["colors"]["navy"])


def _add_context(document: DocumentType, model: dict[str, Any]) -> None:
    # Display the recorded device, host, cables, durations and measurement boundary.
    context = model["test_context"]
    rows = [
        [
            "Device",
            f"{model['identity']['product_name']} | {model['identity']['mxid']}",
            "DepthAI",
            context["depthai_version"],
        ],
        ["Host path", f"{context['host']} | {context['hub']}", "Port", context["port"]],
        [
            "USB 2 cable",
            context["baseline_cable"],
            "Runtime speed",
            model["links"]["usb2"] + " *",
        ],
        [
            "USB 3 cable",
            context["candidate_cable"],
            "Runtime speed",
            model["links"]["usb3"] + " *",
        ],
        [
            "Standard profiles",
            f"{context['standard_duration_s']} s each",
            "Stability repeat",
            f"{context['stability_duration_s']} s",
        ],
        [
            "Power path",
            context["power"],
            "Experimental boundary",
            "Host receipt; no network/browser",
        ],
    ]
    _add_standard_table(
        document,
        ["Field", "Recorded value", "Field", "Recorded value"],
        rows,
        "context",
    )
    _add_note(
        document,
        "INTERPRETATION RULE",
        "A status of complete means evidence collection completed; it does not mean an application or system requirement passed.",
    )


def _add_specifications(document: DocumentType, model: dict[str, Any]) -> None:
    # Display hardware specification rows followed by the three evidence-marker
    # meanings.
    rows = [
        [row["parameter"], row["value"], row["evidence"], row["status"]]
        for row in model["specification_rows"]
    ]
    _add_standard_table(
        document,
        [
            "Parameter",
            "Specification / observed value",
            "Evidence source",
            "Evidence class",
        ],
        rows,
        "specifications",
        status_column=3,
    )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)

    for label, text, color in (
        ("*", "link- or host-dependent", TOKENS["colors"]["teal"]),
        ("†", "scene- or application-dependent; pending test", "8A6500"),
        ("‡", "ambiguous or conflicting supplied source", "9B1C1C"),
    ):
        marker = paragraph.add_run(f"{label} ")
        _set_run_font(
            marker,
            size_pt=9.2,
            color=color,
            bold=True,
        )
        explanation = paragraph.add_run(text + "    ")
        _set_run_font(explanation, size_pt=9.2, color=TOKENS["colors"]["muted"])


def _format_number(value: float | int, decimals: int = 3) -> str:
    if isinstance(value, int):
        return str(value)

    return f"{value:.{decimals}f}"


def _add_comparison(document: DocumentType, model: dict[str, Any]) -> None:
    # Build the side-by-side USB 2/USB 3 stream table with a two-row header.
    table = document.add_table(rows=2 + len(model["comparison_rows"]), cols=11)
    table.style = "Table Grid"
    _apply_table_geometry(table, TABLE_WIDTHS["comparison"])

    # Group the measured columns under separate USB 2 and USB 3 headers.
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Profile"
    table.cell(0, 1).merge(table.cell(1, 1)).text = "Stream"
    table.cell(0, 2).merge(table.cell(0, 5)).text = "USB 2 | HIGH"
    table.cell(0, 6).merge(table.cell(0, 9)).text = "USB 3 | SUPER"
    table.cell(0, 10).merge(table.cell(1, 10)).text = "Observed change *"

    for column, text in zip(range(2, 10), ["FPS", "Gaps", "Mbit/s", "p95 ms"] * 2):
        table.cell(1, column).text = text

    for row_index, row in enumerate(model["comparison_rows"], start=2):
        values = [
            row["profile"],
            row["stream_label"],
            _format_number(row["usb2"]["fps"]),
            row["usb2"]["gaps"],
            _format_number(row["usb2"]["payload_mbps"]),
            _format_number(row["usb2"]["p95_ms"]),
            _format_number(row["usb3"]["fps"]),
            row["usb3"]["gaps"],
            _format_number(row["usb3"]["payload_mbps"]),
            _format_number(row["usb3"]["p95_ms"]),
            row["observation"],
        ]

        for column, value in enumerate(values):
            table.cell(row_index, column).text = str(value)

    _style_table(table, header_rows=2)

    for row_index in range(2, len(table.rows)):
        for column in range(1, 10):
            _format_cell(table.cell(row_index, column), align=WD_ALIGN_PARAGRAPH.CENTER)

    _add_note(
        document,
        "METRIC BOUNDARY",
        "Received FPS counts frames arriving at the host. Sequence gaps identify unobserved sequence numbers but not where loss occurred. Payload Mbit/s excludes USB framing, XLink overhead and data that never reached a host queue. p95 is a software latency estimate in the host-synchronised DepthAI clock domain.",
        fill=TOKENS["colors"]["alternate_fill"],
    )


def _stability_interpretation(measurement: str, usb2: float, usb3: float) -> str:
    # Choose the fixed archived-run interpretation or format a latency/CPU difference.
    if "received FPS" in measurement:
        return "Both observations remained near the configured 30 FPS."

    if "sequence gaps" in measurement:
        return "No sequence gaps were observed in either 60-second run."

    if "p95 latency" in measurement:
        return f"Candidate difference: {usb3 - usb2:+.3f} ms."

    if "temperature" in measurement:
        return "Internal chip telemetry; not the datasheet ambient-air limit."

    return f"Candidate difference: {usb3 - usb2:+.2f} percentage points."


def _add_stability(document: DocumentType, model: dict[str, Any]) -> None:
    # Display stability values and differences, then add the recorded software warning.
    rows = []

    for row in model["stability_rows"]:
        usb2 = row["usb2"]
        usb3 = row["usb3"]
        rows.append(
            [
                row["measurement"],
                _format_number(usb2),
                _format_number(usb3),
                _format_number(usb3 - usb2),
                row["unit"],
                row["status"],
                _stability_interpretation(row["measurement"], usb2, usb3),
            ]
        )

    _add_standard_table(
        document,
        [
            "60-second measurement",
            "USB 2",
            "USB 3",
            "Change",
            "Unit",
            "Status",
            "Interpretation boundary",
        ],
        rows,
        "stability",
        status_column=5,
        numeric_columns=(1, 2, 3),
    )
    warning = (
        model["warnings"][0]
        if model["warnings"]
        else "No runtime warnings were recorded."
    )
    _add_note(
        document,
        "RECORDED SOFTWARE NOTE",
        warning,
        fill=TOKENS["colors"]["pending_fill"],
    )


def _add_constraints(document: DocumentType, model: dict[str, Any]) -> None:
    # Display each constraint, its supporting evidence, consequence and required
    # follow-up.
    rows = [
        [
            f"{row['id']}\n{row['constraint']}",
            row["evidence"],
            row["consequence"],
            row["status"],
            row["follow_up"],
        ]
        for row in model["constraint_rows"]
    ]
    table = _add_standard_table(
        document,
        [
            "Constraint",
            "Supporting evidence",
            "Project consequence",
            "Status",
            "Required follow-up",
        ],
        rows,
        "constraints",
        status_column=3,
    )
    _set_table_density(table, font_size=8.2, vertical_margin_dxa=50)


def _add_boundary_and_evidence(document: DocumentType, model: dict[str, Any]) -> None:
    # Add the supported/pending conclusions and the archived source paths.
    _add_note(
        document,
        "SUPPORTED BY THIS EVIDENCE",
        "The tested OAK-D Lite and cable negotiated SUPER on the MacBook/hub path. USB 3 removed observed sequence gaps from the raw 720p and raw 1080p profiles and restored raw 1080p host delivery to approximately 30 FPS. OAK-side H.264 and StereoDepth completed concurrently for 60 seconds.",
        fill=TOKENS["colors"]["positive_fill"],
    )
    _add_note(
        document,
        "NOT YET ESTABLISHED",
        "Raspberry Pi USB 3 operation, network/browser delivery, overlay quality and cost, ArUco/gauge/valve performance, representative motion/lighting performance, four-second end-to-end delivery, and ten-minute acceptance operation remain pending.",
        fill=TOKENS["colors"]["pending_fill"],
    )
    rows = [
        [reference["label"] + "\n" + reference["id"], reference["path"]]
        for reference in model["evidence_references"]
    ]
    _add_standard_table(
        document,
        ["Evidence record", "Local source path"],
        rows,
        "evidence",
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(
        "Prepared as an evidence reference for EGH455 Group 17. Interpretation and assessment writing remain the student's responsibility."
    )
    _set_run_font(
        run,
        size_pt=9.2,
        color=TOKENS["colors"]["muted"],
        italic=True,
    )


def build_document(model: dict[str, Any]) -> DocumentType:
    """Create the approved compact, presentation-ready Word report."""

    # Validate the model before assembling the six report sections on separate pages.
    validate_report_model(model)
    document = Document()
    _configure_document(document)
    _add_running_furniture(document)
    _add_title_block(document, model)

    _add_section_heading(document, "1. Test Context")
    _add_context(document, model)

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    _add_section_heading(document, "2. Camera Specifications and Observed Hardware")
    _add_specifications(document, model)

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    _add_section_heading(document, "3. USB 2 and USB 3 Benchmark Comparison")
    _add_comparison(document, model)

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    _add_section_heading(document, "4. Stability and System Observations")
    _add_stability(document, model)

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    _add_section_heading(document, "5. Constraints Register")
    _add_constraints(document, model)

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    _add_section_heading(document, "6. Interpretation Boundary and Evidence")
    _add_boundary_and_evidence(document, model)
    return document


def _audit_table(table: Any, table_index: int) -> list[str]:
    # Collect width, indentation, header and cell-span problems without modifying the
    # table.
    issues: list[str] = []
    tbl_pr = table._tbl.tblPr
    table_width = tbl_pr.find(qn("w:tblW"))
    table_indent = tbl_pr.find(qn("w:tblInd"))
    layout = tbl_pr.find(qn("w:tblLayout"))
    grid_widths = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid]
    expected_total = TOKENS["page"]["content_width_dxa"]

    if (
        table_width is None
        or table_width.get(qn("w:type")) != "dxa"
    ):
        issues.append(f"table {table_index}: missing DXA table width")

    elif int(table_width.get(qn("w:w"))) != expected_total:
        issues.append(f"table {table_index}: incorrect table width")

    if (
        table_indent is None
        or int(table_indent.get(qn("w:w"))) != TOKENS["table"]["indent_dxa"]
    ):
        issues.append(f"table {table_index}: incorrect table indent")

    if layout is None or layout.get(qn("w:type")) != "fixed":
        issues.append(f"table {table_index}: layout is not fixed")

    if sum(grid_widths) != expected_total:
        issues.append(f"table {table_index}: grid does not sum to content width")

    if (
        table.rows[0]._tr.find(
            "./w:trPr/w:tblHeader",
            {"w": qn("w:x").split("}")[0][1:]},
        )
        is None
    ):
        issues.append(f"table {table_index}: first row is not repeating header")

    # Check each cell width against the sum of the grid columns covered by its merge.
    for row_index, row in enumerate(table._tbl.findall(qn("w:tr")), start=1):
        grid_index = 0

        for cell in row.findall(qn("w:tc")):
            tc_pr = cell.find(qn("w:tcPr"))

            if tc_pr is None:
                issues.append(
                    f"table {table_index} row {row_index}: cell properties missing"
                )
                continue

            span_element = tc_pr.find(qn("w:gridSpan"))
            span = int(span_element.get(qn("w:val"))) if span_element is not None else 1
            width_element = tc_pr.find(qn("w:tcW"))
            expected_width = sum(grid_widths[grid_index : grid_index + span])

            if (
                width_element is None
                or width_element.get(qn("w:type")) != "dxa"
            ):
                issues.append(
                    f"table {table_index} row {row_index}: missing DXA cell width"
                )

            elif int(width_element.get(qn("w:w"))) != expected_width:
                issues.append(
                    f"table {table_index} row {row_index}: cell width mismatch"
                )

            grid_index += span

        if grid_index != len(grid_widths):
            issues.append(f"table {table_index} row {row_index}: grid span mismatch")

    return issues


def audit_document(document: DocumentType) -> list[str]:
    """Return structural preset or content problems in a generated document."""

    issues: list[str] = []
    section = document.sections[0]
    checks = {
        "page width": (section.page_width.inches, TOKENS["page"]["width_in"]),
        "page height": (section.page_height.inches, TOKENS["page"]["height_in"]),
        "left margin": (section.left_margin.inches, TOKENS["page"]["margin_left_in"]),
        "right margin": (
            section.right_margin.inches,
            TOKENS["page"]["margin_right_in"],
        ),
    }

    for label, (actual, expected) in checks.items():
        if not math.isclose(actual, expected, abs_tol=0.01):
            issues.append(f"{label}: {actual:.3f}, expected {expected:.3f}")

    expected_headings = [
        "1. Test Context",
        "2. Camera Specifications and Observed Hardware",
        "3. USB 2 and USB 3 Benchmark Comparison",
        "4. Stability and System Observations",
        "5. Constraints Register",
        "6. Interpretation Boundary and Evidence",
    ]
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 1"
    ]

    if headings != expected_headings:
        issues.append("Heading 1 structure does not match the approved design")

    if len(document.tables) < 5:
        issues.append("expected at least five report tables")

    forbidden = ("T" + "BD", "T" + "ODO")
    all_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )

    for token in forbidden:
        if token in all_text:
            issues.append(f"placeholder token present: {token}")

    for required in (BASELINE_RUN_ID, CANDIDATE_RUN_ID, "HIGH", "SUPER", "*", "†", "‡"):
        if required not in all_text:
            issues.append(f"required evidence text missing: {required}")

    # Collect table width, header and merged-cell errors from every report table.
    for table_index, table in enumerate(document.tables, start=1):
        issues.extend(_audit_table(table, table_index))

    return issues


def write_report(output_path: Path | None = None) -> Path:
    """Build, structurally audit, and save the report DOCX."""

    project_root = Path(__file__).resolve().parents[1]
    destination = (
        output_path
        or Path(__file__).resolve().parent / "OAK-D-Lite_USB2_USB3_Comparison.docx"
    )
    model = build_report_model(project_root)
    document = build_document(model)

    # Refuse to save the DOCX if its content or table structure fails the audit.
    issues = audit_document(document)

    if issues:
        raise ValueError("document audit failed:\n- " + "\n- ".join(issues))

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def main() -> None:
    print(write_report())


if __name__ == "__main__":
    main()
